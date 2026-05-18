from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt


DEFAULT_INPUT = Path("/Users/max/ucheba/luggify/docs/prediploma_practice_report_2026.md")
DEFAULT_OUTPUT = Path("/Users/max/ucheba/luggify/docs/prediploma_practice_report_2026.docx")


def set_run_font(run, *, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def configure_page(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(14)
        style.font.bold = True


def add_title_line(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    emphasized = (
        text.isupper()
        or text.startswith("«")
        or text.startswith("Тема работы")
        or text in {"ОТЧЕТ", "ПО ПРЕДДИПЛОМНОЙ ПРАКТИКЕ"}
    )
    set_run_font(run, size=14, bold=emphasized)


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph(style=f"Heading {min(level, 3)}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text.strip())
    set_run_font(run, size=14, bold=True)


def add_list_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.left_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text.strip())
    set_run_font(run, size=14, bold=False)


def add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text.strip())
    set_run_font(run, size=14, bold=False)


def render_markdown(source_path: Path, target_path: Path) -> tuple[int, int]:
    text = source_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    document = Document()
    configure_page(document)
    configure_styles(document)

    in_title = False
    word_count = 0
    paragraph_count = 0

    for raw_line in lines:
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped == ":::title":
            in_title = True
            continue

        if stripped == ":::endtitle":
            in_title = False
            continue

        if stripped == ":::pagebreak":
            document.add_page_break()
            continue

        if in_title:
            if not stripped:
                document.add_paragraph()
                continue
            add_title_line(document, stripped)
            word_count += len(stripped.split())
            paragraph_count += 1
            continue

        if not stripped:
            continue

        heading_match = re.match(r"^(#{2,3})\s+(.*)$", stripped)
        if heading_match:
            level = 1 if len(heading_match.group(1)) == 2 else 2
            add_heading(document, heading_match.group(2), level)
            word_count += len(heading_match.group(2).split())
            paragraph_count += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            add_list_paragraph(document, stripped)
            word_count += len(stripped.split())
            paragraph_count += 1
            continue

        add_paragraph(document, stripped)
        word_count += len(stripped.split())
        paragraph_count += 1

    document.save(target_path)
    return word_count, paragraph_count


def main() -> int:
    source_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_INPUT
    target_path = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUTPUT

    if not source_path.exists():
        print(f"Input file not found: {source_path}", file=sys.stderr)
        return 1

    target_path.parent.mkdir(parents=True, exist_ok=True)
    words, paragraphs = render_markdown(source_path, target_path)
    print(f"Generated: {target_path}")
    print(f"Words: {words}")
    print(f"Paragraphs: {paragraphs}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
